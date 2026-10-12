"""
Master PRD Generator for Whisk & Layers
Author: Aditya Nair | Senior Product Manager
Output: Whisk_and_Layers_PRD_Aditya_Nair.docx & Whisk_and_Layers_PRD_Aditya_Nair.pdf
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

from prd_builder_utils import (
    set_cell_background, set_cell_margins, set_cell_borders,
    COLOR_PRIMARY_DARK, COLOR_ACCENT, COLOR_BG_LIGHT, COLOR_BG_ALT,
    COLOR_BORDER, COLOR_BORDER_STRONG, COLOR_TEXT_MAIN, COLOR_TEXT_MUTED,
    RGB_PRIMARY_DARK, RGB_ACCENT, RGB_TEXT_MAIN, RGB_TEXT_MUTED, RGB_WHITE
)

from sections_part1 import build_part1
from sections_part2 import build_part2
from sections_part3 import build_part3
from sections_part4 import build_part4

def build_cover_page(doc, logo_path):
    """Builds a premium, executive cover page."""
    # First section geometry
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.page_width = Inches(8.27)   # A4 Width
    section.page_height = Inches(11.69) # A4 Height
    
    # Enable different first page so cover has no headers/footers
    section.different_first_page_header_footer = True
    
    # Setup Running Header & Footer for Subsequent Pages
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_before = Pt(0)
    hp.paragraph_format.space_after = Pt(4)
    hrun1 = hp.add_run("WHISK & LAYERS  |  ")
    hrun1.font.name = "Segoe UI"
    hrun1.font.size = Pt(8.5)
    hrun1.font.bold = True
    hrun1.font.color.rgb = RGB_PRIMARY_DARK
    
    hrun2 = hp.add_run("Product Requirements Document (PRD) v1.0")
    hrun2.font.name = "Segoe UI"
    hrun2.font.size = Pt(8.5)
    hrun2.font.color.rgb = RGB_TEXT_MUTED

    # Header bottom border line
    hp_pPr = hp._p.get_or_add_pPr()
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="4" w:color="{COLOR_BORDER}"/></w:pBdr>')
    hp_pPr.append(pBdr)

    # Setup Running Footer
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.paragraph_format.space_before = Pt(4)
    fp.paragraph_format.space_after = Pt(0)
    
    # Left part of footer: author info
    frun1 = fp.add_run("Prepared by Aditya Nair | Senior Product Management Portfolio        ")
    frun1.font.name = "Segoe UI"
    frun1.font.size = Pt(8.5)
    frun1.font.color.rgb = RGB_TEXT_MUTED
    
    # Right part of footer: Page X of Y
    frun2 = fp.add_run("Page ")
    frun2.font.name = "Segoe UI"
    frun2.font.size = Pt(8.5)
    frun2.font.color.rgb = RGB_TEXT_MUTED
    
    fld_page = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
    fp._p.append(fld_page)
    
    frun3 = fp.add_run(" of ")
    frun3.font.name = "Segoe UI"
    frun3.font.size = Pt(8.5)
    frun3.font.color.rgb = RGB_TEXT_MUTED
    
    fld_numpages = parse_xml(r'<w:fldSimple %s w:instr="NUMPAGES"/>' % nsdecls('w'))
    fp._p.append(fld_numpages)
    
    # Footer top border line
    fp_pPr = fp._p.get_or_add_pPr()
    fp_pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:top w:val="single" w:sz="6" w:space="4" w:color="{COLOR_BORDER}"/></w:pBdr>')
    fp_pPr.append(fp_pBdr)

    # --- COVER CONTENT ---
    p_top = doc.add_paragraph()
    p_top.paragraph_format.space_before = Pt(24)
    p_top.paragraph_format.space_after = Pt(8)
    
    # Centered Logo
    if os.path.exists(logo_path):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_before = Pt(0)
        p_logo.paragraph_format.space_after = Pt(12)
        run_logo = p_logo.add_run()
        run_logo.add_picture(logo_path, width=Inches(2.1))
    
    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(2)
    run_t = p_title.add_run("WHISK & LAYERS")
    run_t.font.name = "Playfair Display"
    run_t.font.size = Pt(30)
    run_t.font.bold = True
    run_t.font.color.rgb = RGB_PRIMARY_DARK

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(4)
    run_sub = p_sub.add_run("Product Requirements Document (PRD)")
    run_sub.font.name = "Segoe UI"
    run_sub.font.size = Pt(14)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGB_ACCENT

    # Tagline
    p_tag = doc.add_paragraph()
    p_tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tag.paragraph_format.space_before = Pt(0)
    p_tag.paragraph_format.space_after = Pt(18)
    run_tag = p_tag.add_run("India-First Customized Cake Ordering Marketplace Connecting Celebration Seekers with Local Artisan Bakeries")
    run_tag.font.name = "Segoe UI"
    run_tag.font.size = Pt(10)
    run_tag.font.italic = True
    run_tag.font.color.rgb = RGB_TEXT_MAIN

    # Horizontal divider
    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.paragraph_format.space_before = Pt(0)
    p_div.paragraph_format.space_after = Pt(14)
    r_div = p_div.add_run("—" * 38)
    r_div.font.color.rgb = RGBColor(0xCB, 0xD5, 0xE1)

    # Metadata Card Table
    meta_tbl = doc.add_table(rows=10, cols=2)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_tbl.autofit = False

    meta_items = [
        ("Document Identifier", "PRD-WL-2026-V1.0"),
        ("Document Version", "Version 1.0 (Production Specification & Portfolio Baseline)"),
        ("Prepared By", "Aditya Nair"),
        ("Role & Scope", "Senior Product Manager (Strategy, UX Design, TPM, Technical Writing)"),
        ("Target Market", "India (Initial Hubs: Ahmedabad, Vadodara, Rajkot, Surat, Mumbai)"),
        ("Measurement & Currency", "Indian Rupee (INR ₹), Kilogram (0.5 Kg to 3.0 Kg), Indian Mobile (+91)"),
        ("Product Status", "Implemented MVP (Live Web Full-Stack & GitHub Pages Deployment)"),
        ("Release Date", "October 2026"),
        ("Target Audience", "Product Managers, Engineering Leads, Founders, UX Designers, Hiring Committees"),
        ("Live Interactive Demo", "https://desingnerlaunda05.github.io/whisk-and-layers/")
    ]

    for r_idx, (label, val) in enumerate(meta_items):
        row = meta_tbl.rows[r_idx]
        c_lbl = row.cells[0]
        c_val = row.cells[1]
        
        c_lbl.width = Inches(2.3)
        c_val.width = Inches(4.37)
        
        set_cell_background(c_lbl, COLOR_BG_LIGHT if r_idx % 2 == 0 else "FFFFFF")
        set_cell_background(c_val, COLOR_BG_LIGHT if r_idx % 2 == 0 else "FFFFFF")
        set_cell_margins(c_lbl, top=80, bottom=80, left=140, right=140)
        set_cell_margins(c_val, top=80, bottom=80, left=140, right=140)
        
        border = {"val": "single", "sz": "4", "color": COLOR_BORDER}
        set_cell_borders(c_lbl, top=border, bottom=border, left=border, right=border)
        set_cell_borders(c_val, top=border, bottom=border, left=border, right=border)

        p1 = c_lbl.paragraphs[0]
        p1.paragraph_format.space_before = Pt(1)
        p1.paragraph_format.space_after = Pt(1)
        r1 = p1.add_run(label)
        r1.font.name = "Segoe UI"
        r1.font.size = Pt(8.5)
        r1.font.bold = True
        r1.font.color.rgb = RGB_PRIMARY_DARK

        p2 = c_val.paragraphs[0]
        p2.paragraph_format.space_before = Pt(1)
        p2.paragraph_format.space_after = Pt(1)
        r2 = p2.add_run(val)
        r2.font.name = "Segoe UI"
        r2.font.size = Pt(8.5)
        r2.font.color.rgb = RGB_TEXT_MAIN

    p_post = doc.add_paragraph()
    p_post.paragraph_format.space_before = Pt(16)
    p_post.paragraph_format.space_after = Pt(0)
    p_post.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_post = p_post.add_run("CONFIDENTIAL & PROPRIETARY  •  PRODUCT MANAGEMENT PORTFOLIO ARTIFACT")
    r_post.font.name = "Segoe UI"
    r_post.font.size = Pt(8)
    r_post.font.bold = True
    r_post.font.color.rgb = RGB_TEXT_MUTED

    doc.add_page_break()

def build_toc_page(doc):
    """Adds Table of Contents page with valid OpenXML TOC field."""
    h = doc.add_heading(level=1)
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(10)
    run = h.add_run("Table of Contents")
    run.font.name = "Segoe UI"
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = RGB_PRIMARY_DARK

    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.space_before = Pt(0)
    p_intro.paragraph_format.space_after = Pt(8)
    r_in = p_intro.add_run("This document is structured into 41 numbered sections covering strategic product framing, user experience specifications, technical architecture, and execution telemetry.")
    r_in.font.name = "Segoe UI"
    r_in.font.size = Pt(9)
    r_in.font.italic = True
    r_in.font.color.rgb = RGB_TEXT_MUTED

    # Word native TOC field in run
    p_toc = doc.add_paragraph()
    p_toc.paragraph_format.space_before = Pt(4)
    p_toc.paragraph_format.space_after = Pt(12)
    
    run_xml = parse_xml(
        f'<w:r {nsdecls("w")}>'
        f'<w:fldChar w:fldCharType="begin"/>'
        f'<w:instrText xml:space="preserve"> TOC \\o "1-3" \\h \\z \\u </w:instrText>'
        f'<w:fldChar w:fldCharType="separate"/>'
        f'<w:fldChar w:fldCharType="end"/>'
        f'</w:r>'
    )
    p_toc._p.append(run_xml)

    doc.add_page_break()

def main():
    print("=" * 70)
    print("Starting Whisk & Layers Comprehensive PRD Generation...")
    print("Author: Aditya Nair | Senior Product Manager")
    print("=" * 70)

    doc = docx.Document()
    
    # Configure Normal Style typography
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Segoe UI'
    font.size = Pt(9.5)
    font.color.rgb = RGB_TEXT_MAIN

    # Paths
    root_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
    logo_path = os.path.join(root_dir, "assets", "whisk-and-layers-logo.jpeg")
    docx_output_path = os.path.join(root_dir, "Whisk_and_Layers_PRD_Aditya_Nair.docx")
    pdf_output_path = os.path.join(root_dir, "Whisk_and_Layers_PRD_Aditya_Nair.pdf")

    # Step 1: Cover Page
    print("\n[Step 1/6] Building Cover Page...")
    build_cover_page(doc, logo_path)

    # Step 2: Table of Contents
    print("[Step 2/6] Building Table of Contents...")
    build_toc_page(doc)

    # Step 3: Part 1 (Sections 1-10)
    print("[Step 3/6] Compiling Part 1 (Sections 1 to 10)...")
    build_part1(doc)

    # Step 4: Part 2 (Sections 11-20)
    print("[Step 4/6] Compiling Part 2 (Sections 11 to 20)...")
    build_part2(doc)

    # Step 5: Part 3 (Sections 21-30)
    print("[Step 5/6] Compiling Part 3 (Sections 21 to 30)...")
    build_part3(doc)

    # Step 6: Part 4 (Sections 31-41)
    print("[Step 6/6] Compiling Part 4 (Sections 31 to 41)...")
    build_part4(doc)

    # Save DOCX
    print(f"\nSaving Word Document: {docx_output_path}...")
    doc.save(docx_output_path)
    file_size_kb = os.path.getsize(docx_output_path) / 1024
    print(f"DOCX created successfully! File size: {file_size_kb:.1f} KB")

    # Update Fields & Export to PDF via Word COM
    print("\nAutomating Word via COM to refresh fields, calculate pages, and export PDF...")
    try:
        import win32com.client
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        try:
            doc_com = word.Documents.Open(docx_output_path)
            
            # Update fields and TOC
            doc_com.Fields.Update()
            for t in doc_com.TablesOfContents:
                t.Update()
            
            # Compute total pages
            page_count = doc_com.ComputeStatistics(2) # 2 = wdStatisticPages
            print(f">>> Total Rendered Page Count: {page_count} pages <<<")

            # Save updated docx with TOC generated
            doc_com.Save()

            # Export to PDF
            print(f"Exporting PDF to: {pdf_output_path}...")
            doc_com.SaveAs(pdf_output_path, FileFormat=17) # 17 = wdFormatPDF
            pdf_size_kb = os.path.getsize(pdf_output_path) / 1024
            print(f"PDF exported successfully! File size: {pdf_size_kb:.1f} KB")
            
            doc_com.Close()
        finally:
            word.Quit()
    except Exception as e:
        print(f"Note on Word COM automation: {e}")
        print("DOCX deliverable is fully intact and ready to open in Microsoft Word.")

    print("\n" + "=" * 70)
    print("WHISK & LAYERS PRD GENERATION COMPLETE!")
    print(f"Deliverable (DOCX): {docx_output_path}")
    if os.path.exists(pdf_output_path):
        print(f"Deliverable (PDF):  {pdf_output_path}")
    print("=" * 70)

if __name__ == "__main__":
    main()
