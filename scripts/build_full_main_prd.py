# scripts/build_full_main_prd.py
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os
import win32com.client as win32

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="D1D5DB", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_header_footer(doc):
    for s in doc.sections:
        s.different_first_page_header_footer = True
        header = s.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("WHISK & LAYERS  |  PRODUCT REQUIREMENTS DOCUMENT (PRD)")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8)
        hrun.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
        
        footer = s.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        frun1 = fp.add_run("Author: Aditya Nair (Product Portfolio Owner)   |   Candidate MVP Baseline   |   Strict Verification Standard")
        frun1.font.name = "Calibri"
        frun1.font.size = Pt(8)
        frun1.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

def format_run(run, font_name="Calibri", size_pt=10, bold=False, italic=False, color_rgb=(0x33, 0x41, 0x55)):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*color_rgb)

def add_p(doc, text="", style='Normal', space_before=0, space_after=3, line_spacing=1.15, bold=False, italic=False, font_size=9.5, color=(0x33, 0x41, 0x55), align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if text:
        r = p.add_run(text)
        format_run(r, "Calibri", font_size, bold, italic, color)
    return p

def add_h1(doc, text, space_before=11, space_after=3):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    format_run(r, "Calibri", 13.5, bold=True, color_rgb=(0x0F, 0x17, 0x2A))
    return p

def add_h2(doc, text, space_before=7, space_after=2):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    format_run(r, "Calibri", 11, bold=True, color_rgb=(0x9B, 0x2C, 0x2C))
    return p

def add_h3(doc, text, space_before=5, space_after=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    format_run(r, "Calibri", 9.5, bold=True, color_rgb=(0x1E, 0x29, 0x3B))
    return p

def add_bullet(doc, text, bold_prefix="", space_after=2):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.12
    if bold_prefix:
        r_pre = p.add_run(bold_prefix + " ")
        format_run(r_pre, "Calibri", 9.2, bold=True, color_rgb=(0x0F, 0x17, 0x2A))
    r_body = p.add_run(text)
    format_run(r_body, "Calibri", 9.2, color_rgb=(0x33, 0x41, 0x55))
    return p

def add_callout(doc, title, text, border_color="9B2C2C", bg_color="FFF8F6"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(7.0)
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=70, bottom=70, left=140, right=120)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.12
    
    r_title = p.add_run(title + ": ")
    format_run(r_title, "Calibri", 9, bold=True, color_rgb=(0x9B, 0x2C, 0x2C))
    
    r_text = p.add_run(text)
    format_run(r_text, "Calibri", 9, color_rgb=(0x33, 0x41, 0x55))
    
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(2)

def build_table(doc, headers, data, col_widths=None, align_right_cols=None):
    rows = len(data) + 1
    cols = len(headers)
    tbl = doc.add_table(rows=rows, cols=cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    set_table_borders(tbl, color="CBD5E1", sz="4")
    
    # Header Row
    hdr_cells = tbl.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].width = Inches(col_widths[i]) if col_widths else Inches(7.0 / cols)
        set_cell_background(hdr_cells[i], "1E293B")
        set_cell_margins(hdr_cells[i], top=70, bottom=70, left=90, right=90)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if (align_right_cols and i in align_right_cols) else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(title)
        format_run(r, "Calibri", 8.5, bold=True, color_rgb=(0xFF, 0xFF, 0xFF))
    
    # Body Rows
    for r_idx, row_data in enumerate(data):
        row_cells = tbl.rows[r_idx + 1].cells
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].width = Inches(col_widths[c_idx]) if col_widths else Inches(7.0 / cols)
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=55, bottom=55, left=90, right=90)
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if (align_right_cols and c_idx in align_right_cols) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            r = p.add_run(str(val))
            format_run(r, "Calibri", 8.2, color_rgb=(0x1E, 0x29, 0x3B))
            
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(2)
    return tbl

print("Full builder foundation ready.")
